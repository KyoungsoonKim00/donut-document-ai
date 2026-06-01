"""Seed a labeling workbook with PDF filenames.

Writes each PDF filename into column A (starting at A2) of an existing Excel
workbook so a human can fill in the ground-truth fields in the other columns.
"""
from __future__ import annotations

import os
from pathlib import Path

from openpyxl import load_workbook


def write_filenames_to_excel(pdf_dir: str | Path, excel_path: str | Path) -> int:
    """Insert PDF filenames from ``pdf_dir`` into column A of ``excel_path``.

    The workbook must already exist (create it with the label columns first).
    Returns the number of filenames written.
    """
    pdf_dir, excel_path = Path(pdf_dir), Path(excel_path)
    if not excel_path.exists():
        raise FileNotFoundError(
            f"Excel workbook not found: {excel_path}. "
            "Create it with the label column headers first."
        )

    pdf_files = [f for f in sorted(os.listdir(pdf_dir)) if f.lower().endswith(".pdf")]

    wb = load_workbook(excel_path)
    ws = wb.active
    for idx, filename in enumerate(pdf_files, start=2):
        ws[f"A{idx}"] = filename
    wb.save(excel_path)

    print(f"[ok] wrote {len(pdf_files)} filenames to {excel_path}")
    return len(pdf_files)
